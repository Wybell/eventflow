from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path("docs/project/10-EventFlow代码规范与工程质量门禁.docx")
CONTENT_WIDTH = 9360
TABLE_INDENT = 120
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(89, 89, 89)
BLACK = RGBColor(30, 30, 30)


def set_font(run, size=11, color=BLACK, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table_pr = table._tbl.tblPr
    table_width = table_pr.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_pr.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    table_indent = table_pr.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_pr.append(table_indent)
    table_indent.set(qn("w:w"), str(TABLE_INDENT))
    table_indent.set(qn("w:type"), "dxa")
    layout = table_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for column, width in zip(grid.gridCol_lst, widths):
        column.set(qn("w:w"), str(width))

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_width = tc_pr.first_child_found_in("w:tcW")
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_pr.append(tc_width)
            tc_width.set(qn("w:w"), str(widths[index]))
            tc_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.15


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))
    set_font(run, size=8.5, color=GRAY)


def create_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("EventFlow | 工程规范")
    set_font(header_run, size=8.5, color=GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("EventFlow | 代码规范与工程质量门禁 | 第 ")
    set_font(footer_run, size=8.5, color=GRAY)
    add_page_field(footer)
    footer_end = footer.add_run(" 页")
    set_font(footer_end, size=8.5, color=GRAY)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("EventFlow 代码规范与工程质量门禁")
    set_font(title_run, size=22, color=DARK_BLUE, bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run(
        f"第四步：工程初始化 | 规范版本 1.0 | 冻结日期：{date.today().isoformat()}"
    )
    set_font(subtitle_run, size=10.5, color=GRAY)
    return document


def heading(document, text, level=1):
    document.add_paragraph(text, style=f"Heading {level}")


def paragraph(document, text, bold_prefix=None):
    value = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        prefix = value.add_run(bold_prefix)
        set_font(prefix, bold=True)
        rest = value.add_run(text[len(bold_prefix):])
        set_font(rest)
    else:
        run = value.add_run(text)
        set_font(run)
    return value


def bullet_list(document, values):
    for value in values:
        item = document.add_paragraph(style="List Bullet")
        item.paragraph_format.space_after = Pt(4)
        item.paragraph_format.line_spacing = 1.25
        run = item.add_run(value)
        set_font(run)


def number_list(document, values):
    for value in values:
        item = document.add_paragraph(style="List Number")
        item.paragraph_format.space_after = Pt(4)
        item.paragraph_format.line_spacing = 1.25
        run = item.add_run(value)
        set_font(run)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "E8EEF5")
        content = cell.paragraphs[0]
        content.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = content.add_run(header)
        set_font(run, size=9.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            content = cells[index].paragraphs[0]
            run = content.add_run(value)
            set_font(run, size=9.5)
    set_table_geometry(table, widths)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def build_document():
    document = create_document()

    heading(document, "1. 目的、范围与效力")
    paragraph(
        document,
        "本规范是 EventFlow 进入工程初始化后的强制工程基线，覆盖后端、前端、数据库、缓存、消息、版本控制和持续集成。它以既定的模块化单体架构为前提，不引入微服务、分布式事务或尚未被需求证明必要的框架。"
    )
    add_table(document, ["层面", "冻结决策", "执行方式"], [
        ["后端", "Java 17、Spring Boot 3、Spring Security、MyBatis-Plus、Flyway。", "Maven 校验、Spotless、Checkstyle、JUnit 5。"],
        ["前端", "React、TypeScript、Vite、React Router、Zustand、Axios、Ant Design。", "ESLint、Prettier、tsc、Vitest。"],
        ["数据与异步", "MySQL 为事实源；Redis 为名额与缓存；RabbitMQ 只传递异步事件；Outbox 保证可靠投递。", "Flyway、Lua 脚本封装、消息幂等与定时对账。"],
        ["质量门禁", "格式、静态检查、单元测试、集成测试和构建必须在合并前通过。", "受保护主分支与拉取请求审查。"],
    ], [1550, 4300, 3510])
    paragraph(document, "本规范优先级低于已批准的产品规则、接口契约和数据一致性设计；发生冲突时，必须先更新相关 Word 设计文档，再修改代码。")

    heading(document, "2. 总体工程约定")
    add_table(document, ["项目", "强制规则"], [
        ["文本与换行", "源码、配置和脚本统一 UTF-8、LF 结尾、文件末尾保留一个换行；禁止提交 BOM。"],
        ["命名语言", "代码标识、表名、接口路径、事件名和配置键使用英文；面向用户的文案使用中文，必要的技术术语可保留英文。"],
        ["复杂度", "优先清晰而非技巧性写法；一个方法只表达一个用例步骤，嵌套超过 3 层时优先通过提前返回或提取方法降低复杂度。"],
        ["依赖", "新增依赖须说明用途、许可证、维护状态和替代方案；不得为单个简单工具函数引入大型依赖。"],
        ["配置", "环境差异放入配置文件或环境变量；密钥、密码、JWT 私钥和生产连接串不得进入 Git 或日志。"],
        ["注释", "只解释业务约束、状态机原因、并发边界或非常规取舍；禁止重复代码字面含义的注释。TODO 必须带工单号与处理条件。"],
    ], [1850, 7510])

    heading(document, "3. 仓库与模块结构")
    paragraph(document, "工程初始化后采用前后端分目录的模块化单体。后端模块之间通过 application 层接口或领域事件协作，严禁跨模块直接写入对方业务表、Mapper 或 Redis Key。")
    add_table(document, ["目录", "职责与边界"], [
        ["backend/", "Maven Spring Boot 应用；按 auth、activity、sessionquota、registration、waitlist、checkin、notification、admin 等业务模块组织。"],
        ["frontend/", "React 单页应用；按页面功能组织 feature，公共视觉组件和业务组件分别管理。"],
        ["deploy/", "Docker Compose、Nginx、示例环境变量、部署与运维脚本；不得存放真实凭据。"],
        ["docs/", "已批准的项目、接口、数据、测试与交付 Word 文档。文档变更应与实现变更同一拉取请求提交。"],
        ["tools/", "可重复运行的项目辅助脚本；不承载业务运行时逻辑。"],
    ], [1800, 7560])
    add_table(document, ["后端包", "允许内容", "禁止内容"], [
        ["com.eventflow.<module>.api", "Controller、请求/响应 DTO、参数校验、OpenAPI 注解。", "业务状态迁移、SQL、Redis 或 MQ 调用。"],
        ["com.eventflow.<module>.application", "用例服务、事务边界、权限编排、跨模块端口调用。", "HTTP 协议细节、跨模块表直接更新。"],
        ["com.eventflow.<module>.domain", "聚合、枚举、值对象、领域异常、状态转换规则。", "Spring MVC、MyBatis、RedisTemplate 等基础设施依赖。"],
        ["com.eventflow.<module>.infrastructure", "Mapper、持久化实体、Redis Lua 适配器、MQ Publisher/Consumer。", "Controller 返回 DTO、散落的业务判断。"],
        ["com.eventflow.shared", "确有跨模块复用价值的错误码、鉴权上下文、审计、时间和通用基础设施。", "为图省事堆放所有业务对象。"],
    ], [2100, 3550, 3710])

    heading(document, "4. Java 与 Spring Boot 规范")
    heading(document, "4.1 命名、类型与依赖注入", level=2)
    bullet_list(document, [
        "包名全小写，例如 com.eventflow.registration.application；类、接口、枚举和 record 使用 PascalCase；方法、字段和参数使用 camelCase；常量使用 UPPER_SNAKE_CASE。",
        "REST 控制器以 Controller 结尾，应用服务以 Service 结尾，数据访问接口以 Mapper 或 Repository 结尾，转换器以 Mapper 或 Assembler 结尾；不得使用无语义缩写。",
        "请求与响应使用不可变 record；领域值对象优先使用 record；持久化实体可为普通类，但禁止作为 API 请求或响应直接返回。",
        "Spring Bean 统一构造器注入，依赖字段为 final；禁止 @Autowired 字段注入、静态可变单例状态和 Service Locator。",
        "find* 方法可返回 Optional；集合查询一律返回空集合；禁止 Optional 作为字段、参数、序列化 DTO 字段或调用 get()。"
    ])
    heading(document, "4.2 分层、事务与状态机", level=2)
    add_table(document, ["位置", "必须做", "不得做"], [
        ["Controller", "鉴权、@Valid 校验、提取当前用户、调用一个应用用例并返回统一响应。", "编写业务分支、开启事务、返回 Entity。"],
        ["Application Service", "表达用例步骤；明确 @Transactional；调用领域对象完成状态转移；同一事务写事实表与 Outbox。", "在事务内同步等待 RabbitMQ、短信或其他远程调用。"],
        ["Domain", "集中维护报名、候补、活动和签到的合法状态迁移；不变量失败时抛出领域异常。", "依赖 HTTP、SQL、缓存客户端或消息客户端。"],
        ["Infrastructure", "实现数据库、Redis、MQ 和外部服务端口；将技术异常补充上下文后上抛。", "自行决定报名、名额或候补业务状态。"],
    ], [1700, 4050, 3610])
    paragraph(document, "涉及 MySQL 的写操作必须将业务事实变更和 outbox_event 插入放在同一个本地事务中。报名请求的 Redis Lua 原子占位发生在事务之前；若持久化失败，必须通过既定补偿流程释放或核对占位，不能在 Controller 中临时拼接补偿。")
    heading(document, "4.3 异常、响应、日志与安全", level=2)
    bullet_list(document, [
        "领域异常使用 unchecked exception，并带稳定的业务错误码；基础设施异常必须带资源和操作上下文。@RestControllerAdvice 统一转换为既定的 code、message、data、requestId 响应契约。",
        "禁止吞异常、只记录 message 不记录堆栈、用异常完成普通流程控制，或随意捕获 Exception。确有边界捕获时，必须记录上下文并转换或再次抛出。",
        "日志使用 SLF4J 占位符。关键写操作至少带 requestId、actorId、activityId、sessionId、registrationNo、eventId 中的适用字段。",
        "日志、异常信息、审计扩展字段和前端埋点中不得出现密码、JWT、验证码、完整手机号、完整邮箱或未经授权的报名信息；展示与日志均按最小必要原则脱敏。",
        "所有外部输入均用 Bean Validation 校验；组织者和管理员写操作必须同时满足角色权限与组织数据范围校验。"
    ])
    heading(document, "4.4 Java 格式与测试", level=2)
    add_table(document, ["项", "规则"], [
        ["格式", "4 空格缩进、K&R 大括号、单行不超过 120 字符；成员顺序为常量、字段、构造器、公开方法、受保护方法、私有方法。"],
        ["格式化", "Spotless 使用 Palantir Java Format；导入顺序由格式化器统一维护，禁止手动争论空格和换行。"],
        ["静态检查", "Checkstyle 检查命名、行长、禁止通配符导入、空块和不必要的可见性；编译启用全部警告。"],
        ["单元测试", "JUnit 5、AssertJ、Mockito；服务测试不启动 Spring 上下文，不使用 sleep，测试名称描述业务结果。"],
        ["切片与集成测试", "Controller 使用 @WebMvcTest，Mapper 使用 @MybatisTest 或等价切片；Redis、MySQL、RabbitMQ 行为用 Testcontainers 验证关键一致性链路。"],
    ], [1850, 7510])

    heading(document, "5. React 与 TypeScript 规范")
    heading(document, "5.1 文件、组件与状态", level=2)
    add_table(document, ["对象", "规范"], [
        ["文件", "页面、组件、Hook 和 TypeScript 文件使用 kebab-case；React 组件导出名使用 PascalCase；每个文件只表达一个主要职责。"],
        ["组件", "仅使用函数组件与 Hook；Props 显式定义接口或 type；可复用展示组件保持受控和无业务请求副作用。"],
        ["Feature", "页面相关的 API、types、components、hooks 和 tests 放在 feature 内；跨 feature 复用后再提升到 shared。"],
        ["状态", "服务端数据优先由 TanStack Query 管理；Zustand 仅保存跨页面客户端状态，如会话、界面偏好和短暂流程状态。禁止把服务端列表复制到 Zustand。"],
        ["副作用", "请求、订阅和计时器放在自定义 Hook；useEffect 只用于同步外部系统，禁止把可直接计算的派生数据写回 state。"],
    ], [1800, 7560])
    heading(document, "5.2 类型、请求与用户状态", level=2)
    bullet_list(document, [
        "禁止 any；边界不确定数据使用 unknown 并通过运行时校验或明确类型收窄后使用。接口请求与响应类型从 API 契约生成或集中定义，不允许页面内重复声明。",
        "Axios 只通过统一 client 发起，统一注入 Authorization、X-Request-Id、超时、错误映射和 token 刷新策略。组件不得直接拼接基础 URL。",
        "每个异步页面和关键操作必须覆盖 loading、empty、error、success 四种状态；报名、取消、签到等写操作必须防止重复点击，并依据业务 code 给出可恢复提示。",
        "列表 key 必须为稳定业务 ID；不得以数组下标作为会变动列表的 key。表单使用受控字段与统一校验规则，不在提交后静默丢弃服务端字段错误。",
        "正式构建禁止 console.log、console.debug 和未处理 Promise；错误上报只能携带脱敏后的诊断上下文。"
    ])
    heading(document, "5.3 样式、可访问性与前端测试", level=2)
    add_table(document, ["项", "规则"], [
        ["样式", "优先复用已批准的设计令牌、Ant Design 主题变量与局部 CSS Modules；禁止大面积行内 style 和 !important。样式声明顺序：布局、盒模型、文字、视觉、交互。"],
        ["可访问性", "所有图标按钮要有可访问名称与 tooltip；表单控件必须关联 label；键盘可操作；颜色不是状态表达的唯一方式。"],
        ["格式化", "Prettier：2 空格、单引号、分号、尾随逗号 all、printWidth 100；ESLint 负责 React Hooks、未使用变量、Promise 和导入边界。"],
        ["测试", "Vitest + React Testing Library 测试用户可见行为和关键 Hook；关键报名、候补、取消和签到旅程由 Playwright 端到端覆盖。"],
    ], [1850, 7510])

    heading(document, "6. MySQL 与 Flyway 规范")
    paragraph(document, "MySQL 是业务最终事实源。以下规则基于既定的数据一致性设计：Redis 不能替代报名、名额、候补或签到的最终记录，消息队列不能成为数据查询来源。")
    add_table(document, ["对象", "项目规则"], [
        ["字符与时间", "统一 utf8mb4；时间使用 DATETIME(3)，应用统一使用 Asia/Shanghai 业务时区，数据库和应用服务器使用统一时钟策略。"],
        ["表与列", "业务表统一 ef_ 前缀，如 ef_activity、ef_registration；列使用 snake_case；主键为 id BIGINT UNSIGNED AUTO_INCREMENT；关联列为 xxx_id。"],
        ["审计列", "业务事实表必须有 create_time、update_time；create_user_id、update_user_id 只在可追溯的人工操作场景使用。不得机械添加 tenant_id、is_deleted 或 revision。"],
        ["状态", "有生命周期的对象使用 VARCHAR(32) status 并映射枚举；无生命周期的关系表不强加 status。报名状态迁移必须带旧状态条件或等价乐观锁条件。"],
        ["约束", "使用 NOT NULL、默认值、CHECK（可执行时）和唯一索引表达数据不变量；高频报名链路不使用跨模块物理外键，关联列必须有对应索引。"],
        ["索引", "按真实查询设计复合索引，遵循最左前缀；唯一约束优先于代码去重；禁止 SELECT *、无条件大范围更新和未经分页的后台大查询。"],
    ], [1850, 7510])
    bullet_list(document, [
        "所有表结构和数据修正都通过 Flyway。迁移命名为 V<版本>__<英文说明>.sql，例如 V003__create_registration_tables.sql；已在共享环境执行的版本迁移不得修改或删除，只能新增修正迁移。",
        "每个迁移必须有可验证的前置假设、可执行的向前变更和必要的数据回填；大表变更须评估锁表、分批和回滚方案后再合并。",
        "数据库设计变更必须同步更新数据设计 Word 文档和 ER 图，并至少附上受影响查询的索引说明。"
    ])

    heading(document, "7. Redis、RabbitMQ 与 Outbox 调用边界")
    add_table(document, ["组件", "允许调用者", "强制规则"], [
        ["Redis 名额", "sessionquota.infrastructure 的 Lua 适配器与受控补偿任务。", "禁止 Controller、后台管理页或其他模块直接改名额 Key；名额变化必须走原子 Lua 脚本。"],
        ["Redis 缓存", "所属模块的缓存端口实现。", "Key 统一以 ef: 开头并有版本与 TTL；缓存未命中可回源 MySQL，缓存写入失败不能伪造业务成功。"],
        ["RabbitMQ", "Outbox Publisher 与明确的 Consumer。", "业务写事务不直接发布 MQ；生产者必须确认，消费者按 event_id 幂等，失败进入有限重试和死信队列。"],
        ["Outbox", "Application Service 在本地事务写入；Publisher 负责投递与状态。", "事件包含 event_id、event_type、aggregate_type、aggregate_id、payload、occurred_at；消费者不以消息顺序作为正确性前提。"],
        ["补偿与对账", "受控定时任务与管理员受审计操作。", "只能按既定补偿策略修正 Redis、MySQL、Outbox 或队列状态；每次人工补偿必须写 operation_log。"],
    ], [1750, 3100, 4410])
    paragraph(document, "与名额相关的核心不变量保持不变：每个场次 total_quota = available_quota + reserved_quota + confirmed_quota，RESERVED 与 CONFIRMED 均占用名额。任何实现不得绕过该不变量或把候补记录计入已占用名额。")

    heading(document, "8. Git、审查与持续集成门禁")
    add_table(document, ["环节", "强制规则"], [
        ["仓库", "采用单 Git 仓库：frontend、backend、deploy 和 docs 同版本演进。前后端目录隔离不等于拆分仓库；当前阶段不得将其拆为独立仓库。"],
        ["分支", "main 为受保护分支，只接受拉取请求；功能分支使用 feat/<issue>-<short-name>，修复分支使用 fix/<issue>-<short-name>。"],
        ["提交", "使用 Conventional Commits：feat、fix、docs、refactor、test、chore；scope 使用业务模块，例如 feat(registration): add reservation command。一次提交只完成一个可说明的改变。联合功能可在同一 PR 中提交，但前端、后端、迁移、测试和文档应按职责拆成独立提交。"],
        ["拉取请求", "说明目的、风险、测试证据、数据迁移与配置影响；涉及状态机、名额、权限、Redis、MQ 或数据库迁移的改动必须点明回滚与补偿影响。"],
        ["审查", "至少 1 名同级审查者批准；涉及一致性、安全或数据迁移的改动再由模块负责人批准。不得用“后续再补测试”绕过门禁。"],
        ["后端 CI", "mvn spotless:check、checkstyle:check、test、集成测试和 package 全部通过；关键报名链路的失败、重试、幂等和补偿测试不可跳过。"],
        ["前端 CI", "pnpm lint、pnpm format:check、pnpm typecheck、pnpm test、pnpm build 全部通过；关键用户旅程的 Playwright 用例在可用环境执行。"],
    ], [1750, 7610])

    heading(document, "9. 合并前检查清单")
    number_list(document, [
        "需求、接口、状态机和数据设计是否仍与实现一致；若有变化，相关 Word 文档是否已同时更新。",
        "是否跨越了模块边界、绕过了权限校验、Redis Lua、Outbox 或既定补偿流程。",
        "输入校验、错误码、审计、日志字段和敏感信息脱敏是否完整。",
        "数据库迁移是否可重复验证，唯一约束和索引是否覆盖真实读写路径。",
        "格式化、静态检查、类型检查、单元测试、集成测试、构建和必要的端到端测试是否全部通过。",
        "拉取请求是否说明了风险、监控点、回滚方式，以及 Redis、MQ、数据库可能出现的不一致处置。"
    ])

    heading(document, "10. 与参考规范的适配说明")
    paragraph(document, "本规范吸收了知识库的命名、分层、参数校验、日志脱敏、SQL 与索引、React Hook、模块化样式及 Git 审查原则，并根据 EventFlow 当前架构做了以下调整：")
    bullet_list(document, [
        "不采用参考规范中“所有业务表强制 tenant_id、is_deleted、revision”的模板化做法。EventFlow 当前以 organization_id 表达组织数据边界，只在确有业务语义时增加软删除或乐观锁字段。",
        "不采用直接在 Service 中同步发消息的示例。EventFlow 的跨进程异步行为一律经同一 MySQL 事务中的 Outbox 投递。",
        "不采用泛化的 Spring Cloud、Feign、配置中心或微服务规则；当前是模块化单体，这些会增加未被需求证明的运行复杂度。",
        "不采用“手动释放大对象帮助 GC”等规则。Java 运行时由垃圾回收器管理，代码应避免不必要的大对象保留和静态可变缓存。"
    ])
    paragraph(document, "本文件完成后，第四步的“代码规范冻结”子步骤结束。下一子步骤才是创建实际工程骨架、构建配置与质量门禁文件；在收到明确指令前，不创建业务代码。")

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
