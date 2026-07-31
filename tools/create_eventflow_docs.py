from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path("docs/project")
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(89, 89, 89)
CONTENT_WIDTH_DXA = 9360


def set_font(run, name="Calibri", size=11, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_geometry(table, widths):
    table.autofit = False
    table_pr = table._tbl.tblPr
    table_width = table_pr.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_pr.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    layout = table_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.15


def create_document(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_paragraph = doc.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(3)
    title_run = title_paragraph.add_run(title)
    set_font(title_run, size=22, color=DARK_BLUE, bold=True)
    subtitle_paragraph = doc.add_paragraph()
    subtitle_paragraph.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle_paragraph.add_run(subtitle)
    set_font(subtitle_run, size=10.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("EventFlow | " + str(date.today()))
    set_font(footer_run, size=8.5, color=GRAY)
    return doc


def add_heading(doc, text, level=1):
    doc.add_paragraph(text, style=f"Heading {level}")


def add_paragraph(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        set_font(prefix, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_font(rest)
    else:
        run = paragraph.add_run(text)
        set_font(run)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        set_font(run)


def add_steps(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        set_font(run)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        shade(header_cells[index], "E8EEF5")
        paragraph = header_cells[index].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_font(run, size=9.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            paragraph = cells[index].paragraphs[0]
            run = paragraph.add_run(value)
            set_font(run, size=9.5)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def save(doc, filename):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DIR / filename)


def product_requirements():
    doc = create_document("EventFlow 活动报名与名额预约平台", "产品需求文档（MVP） | 版本 1.0 | 状态：讨论中")
    add_heading(doc, "1. 项目定位")
    add_paragraph(doc, "EventFlow 面向企业、技术社区和培训机构，提供统一的活动、场次、名额和报名模型。首版以技术沙龙和企业培训为示例，系统能力保持通用。")
    add_heading(doc, "2. 目标与边界")
    add_table(doc, ["类别", "定义"], [
        ["业务目标", "让组织者可发布多场次活动，并让用户完成可追踪、不可超卖的报名与签到。"],
        ["技术目标", "在高并发报名时通过 Redis 原子名额控制和异步落库，保障最终一致性。"],
        ["本期范围", "创建与发布活动、场次和名额、报名确认、取消、超时回收、候补递补、二维码签到、后台查询。"],
        ["本期不做", "真实支付、复杂座位图、多租户、分库分表、微服务拆分、推荐系统、真实短信和复杂动态表单。"],
    ], [2400, 6960])
    add_heading(doc, "3. 用户角色与权限")
    add_table(doc, ["角色", "核心权限"], [
        ["普通用户", "浏览活动；选择场次；提交、确认或取消报名；查看候补状态、签到二维码和报名记录；现场签到。"],
        ["活动组织者", "创建草稿；配置场次、名额和报名时间；发布或下线活动；查看名单、审核报名与查看签到。"],
        ["平台管理员", "管理用户与组织者；审核活动；查看平台数据、操作日志；处理异常报名与名额补偿。"],
    ], [1800, 7560])
    add_heading(doc, "4. MVP 功能清单")
    add_table(doc, ["优先级", "模块", "验收结果"], [
        ["P0", "活动与场次", "组织者可创建草稿、添加多个场次与名额、设置报名窗口并发布。"],
        ["P0", "报名与名额", "报名不超卖、不重复；用户取得临时保留名额后可确认或取消。"],
        ["P0", "超时与候补", "未确认保留在 10 分钟后释放；释放后按候补顺序分配临时名额。"],
        ["P0", "后台与审计", "组织者可查看报名及签到；管理员可追溯关键操作。"],
        ["P1", "通知与统计", "提供站内通知、活动报名数据与签到统计。"],
    ], [1050, 1800, 6510])
    add_heading(doc, "5. 关键业务规则")
    add_bullets(doc, [
        "服务器是报名时间和剩余名额的唯一判定方，前端倒计时仅用于展示。",
        "同一用户在同一场次最多存在一条有效报名或候补记录；数据库唯一索引作为最终兜底。",
        "名额锁定状态为 RESERVED，默认有效期 10 分钟；到期后转换为 EXPIRED 并回补名额。",
        "名额满时按进入时间进入 WAITING 队列；释放名额后仅向最早候补用户发放一个临时保留资格。",
        "活动开始后仅 CONFIRMED 报名可以签到，签到成功后状态转为 CHECKED_IN。",
    ])
    add_heading(doc, "6. 成功指标")
    add_bullets(doc, [
        "报名主链路 P0 功能全部可验收，异常路径可复现并有补偿手段。",
        "压测下不存在超卖或重复报名；名额与报名记录可完成对账。",
        "关键管理操作与状态变化均保留可查询的审计记录。",
    ])
    save(doc, "01-EventFlow产品需求文档.docx")


def workflows():
    doc = create_document("EventFlow 核心流程与状态机", "业务流程、状态定义与异常处理 | 版本 1.0")
    add_heading(doc, "1. 活动生命周期")
    add_table(doc, ["状态", "进入条件", "允许操作"], [
        ["DRAFT", "组织者新建活动", "编辑活动、场次、名额与报名窗口。"],
        ["PUBLISHED", "通过发布校验", "对用户可见，等待报名窗口开始。"],
        ["REGISTRATION_OPEN", "当前时间进入报名窗口", "接受报名、候补、取消与名额回补。"],
        ["REGISTRATION_CLOSED", "报名截止或管理员关闭", "停止新报名，保留已确认报名与签到准备。"],
        ["ENDED", "活动结束", "保留查询与报表，不再修改报名。"],
        ["OFFLINE", "管理员或组织者下线", "停止对外访问，保留审计与历史数据。"],
    ], [2100, 3300, 3960])
    add_heading(doc, "2. 用户报名主流程")
    add_steps(doc, [
        "用户选择活动场次并提交包含 requestId 的报名请求。",
        "服务端校验活动状态、报名时间、用户资格和请求幂等性。",
        "Redis Lua 脚本原子检查剩余名额、重复占位并写入临时保留。",
        "名额充足时生成 registrationNo，投递报名事件；不足时写入候补队列。",
        "消费者幂等创建报名记录；成功后向用户展示处理中或保留成功状态。",
        "用户在保留期内确认，状态转为 CONFIRMED；未确认则由延迟消息触发过期回收。",
    ])
    add_heading(doc, "3. 报名状态机")
    add_table(doc, ["当前状态", "可转移状态", "触发条件"], [
        ["RESERVED", "CONFIRMED / CANCELLED / EXPIRED", "用户确认、主动取消或保留到期。"],
        ["WAITING", "RESERVED / CANCELLED", "候补递补成功或用户主动退出。"],
        ["CONFIRMED", "CANCELLED / CHECKED_IN", "用户在规则允许期取消或活动现场签到。"],
        ["CANCELLED", "终态", "名额已释放；需根据候补规则触发递补。"],
        ["EXPIRED", "终态", "锁定已过期，名额已回补。"],
        ["CHECKED_IN", "终态", "签到成功，保留时间与操作人。"],
    ], [2100, 3150, 4110])
    add_heading(doc, "4. 候补递补流程")
    add_steps(doc, [
        "取消报名或超时回收释放名额，并记录可递补事件。",
        "读取该场次最早的有效候补记录，使用原子操作获取递补资格。",
        "将候补记录置为 RESERVED，写入新的保留到期时间，并发送通知。",
        "候补用户确认后转为 CONFIRMED；再次过期或取消则继续下一位候补。",
    ])
    add_heading(doc, "5. 异常处理原则")
    add_bullets(doc, [
        "MQ 消费失败需有限次重试；超过次数进入死信队列并由管理员或补偿任务处理。",
        "Redis 已占位但数据库落库失败时，补偿任务必须将报名补写成功或回补名额。",
        "状态转移必须带条件更新，禁止通过普通更新覆盖终态或逆向回退。",
        "所有人工补偿、管理员关闭和名额调整均写入操作日志。",
    ])
    save(doc, "02-EventFlow核心流程与状态机.docx")


def architecture_data():
    doc = create_document("EventFlow 技术架构与数据设计", "模块化单体架构 | Java 17 / Spring Boot 3 / React")
    add_heading(doc, "1. 架构决策")
    add_paragraph(doc, "采用模块化单体而不是微服务。当前系统规模以清晰的模块边界、单一部署单元和事件机制换取更低的部署、调试与分布式事务成本，同时保留未来独立扩缩容时的拆分空间。")
    add_heading(doc, "2. 模块边界")
    add_table(doc, ["模块", "职责", "禁止事项"], [
        ["auth", "认证、JWT、用户和角色权限。", "不承载活动或报名业务规则。"],
        ["activity / session", "活动草稿、发布、场次与报名窗口。", "不直接修改报名状态。"],
        ["quota / registration", "名额占位、报名状态机、幂等、取消与确认。", "不绕过 Redis 原子操作直接扣减库存。"],
        ["waitlist / notification", "候补顺序、递补通知、提醒事件。", "不直接篡改其他模块业务表。"],
        ["checkin / admin", "二维码签到、运营查询、审计和人工补偿。", "不绕过授权执行管理动作。"],
    ], [1900, 4400, 3060])
    add_heading(doc, "3. 一致性方案")
    add_steps(doc, [
        "Redis Lua 完成资格校验、重复占位检查、名额扣减和保留过期时间写入。",
        "RabbitMQ 异步传递报名、超时回收、候补递补、通知等领域事件。",
        "消费者以 requestId 和数据库唯一索引实现幂等，并在事务内写入业务记录。",
        "推荐增加 outbox_event 表：本地事务写入报名与待投递事件，由后台任务可靠投递 MQ。",
        "定时任务对账 Redis 占位、报名记录和事件状态，补写或回补异常数据。",
    ])
    add_heading(doc, "4. 核心数据表")
    add_table(doc, ["表", "关键字段与约束"], [
        ["user", "id、手机号或邮箱、角色、状态、创建与更新时间。"],
        ["activity", "id、组织者、标题、状态、报名开始与结束时间、发布信息。"],
        ["activity_session", "id、activity_id、地点、开始结束时间、场次状态。"],
        ["activity_quota", "session_id、capacity、available_quota、reserved_quota、version。"],
        ["registration", "user_id、session_id、registration_no、status、reservation_expire_time、确认和取消时间；唯一索引 UNIQUE(user_id, session_id)。"],
        ["waitlist", "user_id、session_id、status、queue_no、reserved_until；唯一索引 UNIQUE(user_id, session_id)。"],
        ["check_in / operation_log / outbox_event", "签到时间、操作主体、业务对象、事件载荷、投递与消费状态。"],
    ], [2300, 7060])
    add_heading(doc, "5. Redis 与消息设计")
    add_table(doc, ["类型", "键或事件", "用途"], [
        ["Redis", "activity:quota:{sessionId}", "场次可用名额。"],
        ["Redis", "registration:idempotent:{userId}:{sessionId}:{requestId}", "重复请求拦截与结果复用。"],
        ["Redis", "activity:waitlist:{sessionId}", "候补队列。"],
        ["MQ", "registration.reserved / registration.confirmed", "异步落库、通知和后续处理。"],
        ["MQ", "registration.expire / quota.released", "延迟回收与候补递补。"],
    ], [2650, 4300, 2410])
    add_heading(doc, "6. 技术栈")
    add_bullets(doc, [
        "后端：Java 17、Spring Boot 3、Spring Security + JWT、MyBatis-Plus、Flyway、OpenAPI。",
        "基础设施：MySQL 8、Redis、RabbitMQ、MinIO（用于活动素材或二维码文件时启用）。",
        "前端：React、TypeScript、Vite、React Router、Zustand、Axios、Ant Design；TanStack Query 可选。",
        "部署与验证：Docker Compose、Nginx、JMeter 或 Apache Bench、GitHub Actions 或 Jenkins。",
    ])
    save(doc, "03-EventFlow技术架构与数据设计.docx")


def testing_acceptance():
    doc = create_document("EventFlow 测试与验收方案", "功能、并发、一致性与上线验收 | 版本 1.0")
    add_heading(doc, "1. 测试范围")
    add_table(doc, ["类别", "必须覆盖的风险"], [
        ["功能测试", "活动发布、报名、确认、取消、候补、签到、权限和后台查询。"],
        ["接口测试", "参数校验、鉴权、幂等 requestId、状态冲突、重复提交和错误码。"],
        ["并发测试", "多个用户抢同一场次；同一用户并发提交；名额耗尽与候补入队。"],
        ["一致性测试", "Redis 预扣、MQ 落库、失败重试、死信和补偿任务对账。"],
        ["安全与可靠性", "JWT 过期、越权、敏感字段脱敏、日志审计、重启后任务恢复。"],
    ], [2100, 7260])
    add_heading(doc, "2. P0 验收用例")
    add_table(doc, ["编号", "场景", "预期结果"], [
        ["REG-001", "100 个名额下发起超过 100 个并发报名", "成功确认和保留总数不超过 100；其余请求进入候补或返回满额。"],
        ["REG-002", "同一用户以相同或不同 requestId 重复提交", "只生成一条有效报名或候补记录；重复请求返回原处理结果。"],
        ["REG-003", "用户保留后 10 分钟未确认", "报名转为 EXPIRED，Redis 名额回补，必要时递补候补用户。"],
        ["REG-004", "已确认用户取消报名", "报名转为 CANCELLED，名额释放，并按顺序通知候补用户。"],
        ["REG-005", "MQ 消费失败并重试", "不会重复落库；重试耗尽后进入死信并可由补偿任务处理。"],
        ["REG-006", "扫码签到与重复扫码", "仅 CONFIRMED 记录可签到；重复扫码不产生第二条签到记录。"],
    ], [1000, 3310, 5050])
    add_heading(doc, "3. 压测策略")
    add_bullets(doc, [
        "准备固定场次、固定名额和独立压测用户；压测环境不得连接生产数据库或消息队列。",
        "分别测试单用户重复点击、多用户并发报名、热点场次满额和 MQ 消费积压。",
        "记录吞吐量、P95/P99 响应时间、错误率、Redis/MQ/MySQL 指标、名额差异与补偿次数。",
        "压测结束后执行对账：有效报名数 + 有效保留数 + 可用名额应等于场次总容量。",
    ])
    add_heading(doc, "4. 上线验收门槛")
    add_bullets(doc, [
        "P0 用例全部通过，无阻塞级缺陷；关键缺陷均有验证过的修复记录。",
        "并发报名不发生超卖和重复报名；异常注入后可通过重试、死信或补偿恢复。",
        "备份、回滚、日志、监控和告警方式已验证；管理员可查询关键操作和异常事件。",
        "前端在桌面与移动浏览器中完成主链路验证，含 loading、错误、空状态与防重复点击。",
    ])
    save(doc, "04-EventFlow测试与验收方案.docx")


def deployment_delivery():
    doc = create_document("EventFlow 部署与交付方案", "Docker Compose 与 Nginx 部署基线 | 版本 1.0")
    add_heading(doc, "1. 部署拓扑")
    add_table(doc, ["组件", "职责", "生产要求"], [
        ["eventflow-web", "React 静态资源与反向代理入口。", "Nginx 提供 HTTPS、SPA 回退、静态缓存与 /api 代理。"],
        ["eventflow-api", "Spring Boot 模块化单体应用。", "通过环境变量读取连接信息；提供健康检查与结构化日志。"],
        ["mysql", "业务数据、Flyway 版本与 outbox 事件。", "持久化数据卷、最小权限账号、定期备份与恢复演练。"],
        ["redis", "名额、幂等、候补队列和缓存。", "持久化策略、内存阈值、访问认证与监控。"],
        ["rabbitmq", "报名、过期、通知与补偿事件。", "持久化队列、死信队列、管理告警与消费者监控。"],
        ["minio（可选）", "活动封面、二维码等对象存储。", "桶权限隔离、生命周期策略与备份。"],
    ], [1900, 3670, 3790])
    add_heading(doc, "2. 环境配置原则")
    add_bullets(doc, [
        "禁止将密码、JWT 密钥、MQ 账号、对象存储密钥和域名写入仓库；通过 .env 或部署平台密钥注入。",
        "本地、测试、生产使用独立数据库、Redis、MQ 虚拟主机和存储桶；禁止跨环境连接。",
        "Flyway 是唯一的数据库结构变更入口；生产禁止手工修改表结构。",
        "服务启动时校验配置完整性，健康检查需覆盖 MySQL、Redis 和 RabbitMQ 连接状态。",
    ])
    add_heading(doc, "3. 发布流程")
    add_steps(doc, [
        "合并经过代码评审和自动化测试的版本，生成可追溯的 Git 标签与构建产物。",
        "在测试环境执行 Flyway 迁移、服务部署和冒烟测试，确认报名主链路与监控数据。",
        "备份生产数据库和当前镜像版本，记录本次发布版本、变更项与回滚版本。",
        "按顺序执行数据库迁移、后端镜像更新、前端静态资源发布与 Nginx reload。",
        "验证健康检查、登录、活动查询、报名、MQ 消费和管理员审计日志。",
        "异常时按预案回滚应用镜像；涉及数据迁移时使用预先演练的向后兼容或回滚脚本。",
    ])
    add_heading(doc, "4. 监控与告警")
    add_table(doc, ["指标", "告警条件", "处置方向"], [
        ["报名失败率", "短时间超过阈值", "检查 API 错误、Redis Lua 返回码和 MQ 状态。"],
        ["名额对账差异", "场次容量不守恒", "暂停异常场次，运行补偿并审计差异。"],
        ["MQ 堆积或死信", "消息持续增长或出现死信", "检查消费者、重试原因与幂等记录。"],
        ["服务健康检查", "依赖不可用或响应超时", "按依赖类型降级、恢复或回滚。"],
    ], [2200, 3000, 4160])
    add_heading(doc, "5. 交付清单")
    add_bullets(doc, [
        "前后端源码、Dockerfile、docker-compose 配置、Nginx 配置样例和环境变量模板。",
        "Flyway 迁移脚本、初始化示例数据、OpenAPI 文档与管理员账号初始化说明。",
        "测试报告、压测报告、发布记录、回滚方案、监控告警说明和操作手册。",
    ])
    save(doc, "05-EventFlow部署与交付方案.docx")


def main():
    product_requirements()
    workflows()
    architecture_data()
    testing_acceptance()
    deployment_delivery()


if __name__ == "__main__":
    main()
