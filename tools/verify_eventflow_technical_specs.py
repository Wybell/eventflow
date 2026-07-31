from pathlib import Path

from docx import Document


ROOT = Path("docs/project")
CHECKS = {
    "07-EventFlow详细技术架构与模块设计.docx": ["模块化单体", "Outbox", "X-Request-Id"],
    "08-EventFlow数据Redis与消息一致性设计.docx": ["total_quota = available_quota + reserved_quota + confirmed_quota", "Redis Lua", "RabbitMQ"],
    "09-EventFlow接口与权限契约设计.docx": ["X-Request-Id", "REGISTRATION_RESERVED", "资源归属"],
}


def main() -> int:
    valid = True
    for filename, expected_texts in CHECKS.items():
        document = Document(ROOT / filename)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        content = "\n".join(paragraphs + table_cells)
        missing = [expected_text for expected_text in expected_texts if expected_text not in content]
        matched = not missing
        print(f"{filename}: {matched}")
        if missing:
            print(f"Missing: {', '.join(missing)}")
        valid = valid and matched
    return 0 if valid else 1


if __name__ == "__main__":
    raise SystemExit(main())
