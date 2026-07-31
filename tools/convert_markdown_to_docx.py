from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(89, 89, 89)
CONTENT_WIDTH_DXA = 9360


def set_font(run, name: str, size: float, color: RGBColor | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def configure_table(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    set_font(run, "Calibri", 9.5)


def add_text_paragraph(doc: Document, text: str, style: str | None = None, code: bool = False) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    if code:
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(text)
        set_font(run, "Consolas", 9)
        return

    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", 10)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, "Calibri", 11, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, "Calibri", 11)


def add_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        cells = [item.strip() for item in line.strip().strip("|").split("|")]
        if cells and not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    widths = [CONTENT_WIDTH_DXA // columns] * columns
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            value = row[column_index] if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value.replace("`", ""))
            set_font(run, "Calibri", 9.5, bold=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, "E8EEF5")
    configure_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = doc.styles.add_style("EventFlowTitle", WD_STYLE_TYPE.PARAGRAPH)
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(22)
    title.font.color.rgb = DARK_BLUE
    title.font.bold = True
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("EventFlow Project Reference")
    set_font(footer_run, "Calibri", 8.5, GRAY)


def convert_markdown(source: Path, destination: Path) -> None:
    lines = source.read_text(encoding="utf-8-sig").splitlines()
    doc = Document()
    configure_document(doc)
    title_written = False
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.strip().startswith("```"):
            if in_code:
                for code_line in code_lines:
                    add_text_paragraph(doc, code_line, code=True)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.lstrip().startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if not title_written and level == 1:
                paragraph = doc.add_paragraph(style="EventFlowTitle")
                run = paragraph.add_run(text)
                set_font(run, "Calibri", 22, DARK_BLUE, True)
                title_written = True
            else:
                doc.add_paragraph(text, style=f"Heading {min(level, 3)}")
            index += 1
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if bullet:
            add_text_paragraph(doc, bullet.group(1), style="List Bullet")
            index += 1
            continue
        ordered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if ordered:
            add_text_paragraph(doc, ordered.group(1), style="List Number")
            index += 1
            continue
        if re.fullmatch(r"[-*_]{3,}", line.strip()):
            index += 1
            continue
        add_text_paragraph(doc, line.strip())
        index += 1
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: convert_markdown_to_docx.py SOURCE_DIR DESTINATION_DIR")
        return 2
    source_root = Path(sys.argv[1])
    destination_root = Path(sys.argv[2])
    for source in source_root.rglob("*.md"):
        relative = source.relative_to(source_root).with_suffix(".docx")
        convert_markdown(source, destination_root / relative)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
