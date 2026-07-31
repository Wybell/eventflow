from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path("docs/project")
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(89, 89, 89)


def set_font(run, size=11, color=None, bold=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_width(cell, width):
    properties = cell._tc.get_or_add_tcPr()
    value = properties.first_child_found_in("w:tcW")
    if value is None:
        value = OxmlElement("w:tcW")
        properties.append(value)
    value.set(qn("w:w"), str(width))
    value.set(qn("w:type"), "dxa")


def create_document(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_paragraph = doc.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(3)
    title_run = title_paragraph.add_run(title)
    set_font(title_run, 22, DARK_BLUE, True)
    subtitle_paragraph = doc.add_paragraph()
    subtitle_paragraph.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle_paragraph.add_run(subtitle)
    set_font(subtitle_run, 10.5, GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("EventFlow Technical Design")
    set_font(footer_run, 8.5, GRAY)
    return doc


def heading(doc, text, level=1):
    doc.add_paragraph(text, style=f"Heading {level}")


def paragraph(doc, text):
    item = doc.add_paragraph()
    run = item.add_run(text)
    set_font(run)


def bullets(doc, values):
    for value in values:
        item = doc.add_paragraph(style="List Bullet")
        item.paragraph_format.space_after = Pt(4)
        run = item.add_run(value)
        set_font(run)


def steps(doc, values):
    for value in values:
        item = doc.add_paragraph(style="List Number")
        item.paragraph_format.space_after = Pt(4)
        run = item.add_run(value)
        set_font(run)


def table(doc, headers, rows, widths):
    value = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        cell = value.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "E8EEF5")
        item = cell.paragraphs[0]
        item.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = item.add_run(header)
        set_font(run, 9.5, bold=True)
    for row in rows:
        cells = value.add_row().cells
        for index, text in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            item = cell.paragraphs[0]
            run = item.add_run(text)
            set_font(run, 9.5)
    for row in value.rows:
        for index, cell in enumerate(row.cells):
            set_width(cell, widths[index])
            for item in cell.paragraphs:
                item.paragraph_format.space_after = Pt(3)
                item.paragraph_format.line_spacing = 1.15
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def save(doc, filename):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DIR / filename)


def architecture_spec():
    doc = create_document("EventFlow 详细技术架构与模块设计", "模块化单体 | Java 17 + Spring Boot 3 | 技术设计版本 1.0")
    heading(doc, "1. 架构目标")
    paragraph(doc, "系统采用模块化单体：一次部署、单一数据源，但通过清晰的包边界、服务接口与领域事件避免业务模块相互穿透。此选择适合当前 MVP，降低部署和分布式事务成本，同时为未来独立扩缩容保留拆分路径。")
    table(doc, ["质量目标", "设计约束"], [
        ["高并发", "报名请求先进入 Redis Lua 原子链路，避免热点名额竞争直接压向 MySQL。"],
        ["最终一致", "使用 Outbox、RabbitMQ、幂等消费、死信和定时对账，不宣称 Redis 与 MySQL 强一致。"],
        ["可维护", "Controller 只处理协议；Service 承载用例；模块不直接更新其他模块业务表。"],
        ["可观测", "关键状态迁移、消息投递、补偿与管理员操作必须具备 requestId、registrationNo 或 businessId 追踪。"],
    ], [2200, 7160])
    heading(doc, "2. 运行拓扑")
    steps(doc, [
        "React Web 通过 Nginx 提供静态资源、HTTPS、SPA 路由回退和 /api 反向代理。",
        "Spring Boot 应用处理 REST API、认证授权、业务事务、Redis 与 RabbitMQ 交互。",
        "MySQL 保存最终业务事实、Flyway 迁移、Outbox 事件和操作审计。",
        "Redis 保存高频名额、幂等占位、候补队列与活动详情缓存。",
        "RabbitMQ 传递报名、超时、候补、通知和补偿事件；MinIO 仅在活动封面或二维码文件需要持久化时启用。",
    ])
    heading(doc, "3. 后端模块边界")
    table(doc, ["模块", "职责", "公开能力", "不允许的耦合"], [
        ["auth", "用户、角色、登录、JWT、权限判定。", "当前用户与授权检查。", "不处理活动或报名状态。"],
        ["activity", "活动草稿、发布、下线、报名窗口。", "活动可报名状态查询。", "不直接创建报名记录。"],
        ["session-quota", "场次、总容量、Redis 名额加载与回补。", "原子占位和名额释放。", "不写用户个人信息。"],
        ["registration", "报名状态机、确认、取消、幂等与 Outbox。", "报名创建、查询和状态转移。", "不直接操作候补排序。"],
        ["waitlist", "候补入队、递补资格和保留到期。", "领取最早候补资格。", "不跳过 registration 状态机。"],
        ["checkin", "动态二维码校验、现场签到、签到记录。", "签到和签到查询。", "不修改活动名额。"],
        ["notification-admin", "站内通知、平台治理、审计、人工补偿。", "通知投递、后台查询与补偿。", "不绕过权限和操作日志。"],
    ], [1550, 2450, 2550, 2810])
    heading(doc, "4. Spring Boot 包结构")
    table(doc, ["目录", "规则"], [
        ["com.eventflow.<module>.controller", "仅处理请求映射、Bean Validation、鉴权和 DTO 转换；Controller 命名为 *Controller。"],
        ["com.eventflow.<module>.application", "按用例组织 Service；使用构造器注入；事务边界在应用服务中明确。"],
        ["com.eventflow.<module>.domain", "实体、枚举、领域服务、领域事件和领域异常；状态转换集中在此处。"],
        ["com.eventflow.<module>.infrastructure", "MyBatis-Plus Mapper、Redis、MQ、对象存储和外部实现。"],
        ["com.eventflow.<module>.dto", "请求与响应 DTO 优先使用 Java record；不得把 Entity 直接返回给 API。"],
        ["com.eventflow.shared", "统一错误响应、认证上下文、时钟、序列号、审计与基础配置。"],
    ], [3150, 6210])
    heading(doc, "5. 编码与事务规则")
    bullets(doc, [
        "Service 采用构造器注入和 final 依赖，不使用 @Autowired 字段注入。",
        "输入使用 Bean Validation；find* 方法返回 Optional；领域错误使用明确的 unchecked exception 并由 @RestControllerAdvice 统一映射。",
        "数据库事务仅覆盖 MySQL 内的状态更新和 Outbox 写入，不在事务中同步调用 RabbitMQ 或远程服务。",
        "状态更新必须使用条件更新，例如 WHERE id = ? AND status = ?，防止并发覆盖与非法逆向迁移。",
        "日志使用结构化字段：requestId、userId、activityId、sessionId、registrationNo、eventId；禁止记录密码、JWT 或完整手机号。",
    ])
    heading(doc, "6. 请求生命周期")
    steps(doc, [
        "网关或过滤器生成/读取 X-Request-Id，并写入日志上下文。",
        "Spring Security 校验 JWT，构建当前用户和角色上下文。",
        "Controller 校验请求 DTO 后调用模块 Application Service。",
        "报名服务调用 Redis Lua 获取占位结果；结果为成功时创建报名记录和 Outbox 事件。",
        "Outbox Publisher 异步投递 RabbitMQ，消费者按 eventId 幂等处理通知、超时或候补后续动作。",
        "响应统一返回 data、code、message、requestId；前端以业务 code 区分候补、重复提交和保留已过期。",
    ])
    heading(doc, "7. 安全与演进")
    bullets(doc, [
        "JWT 设置短访问令牌和可撤销刷新令牌；接口采用基于角色和资源归属的双重授权。",
        "组织者只能读取和操作自己组织内的活动；管理员可跨组织治理并留下审计记录。",
        "未来拆分优先级：notification、checkin 或高峰 registration 仅在出现独立发布、独立扩容或故障隔离证据后拆出。",
        "不在当前阶段引入 Nacos、Feign、分布式事务框架或 Kubernetes；这些不会解决 MVP 的主要业务风险。",
    ])
    save(doc, "07-EventFlow详细技术架构与模块设计.docx")


def data_event_spec():
    doc = create_document("EventFlow 数据、Redis 与消息一致性设计", "MySQL 8 + Redis + RabbitMQ + Outbox | 技术设计版本 1.0")
    heading(doc, "1. 数据设计原则")
    bullets(doc, [
        "MySQL 是最终业务事实来源；Redis 是短生命周期并发控制与缓存；RabbitMQ 是异步事件通道。",
        "所有业务表使用 BIGINT UNSIGNED AUTO_INCREMENT 主键、utf8mb4、create_time 和 update_time；删除采用业务状态而非物理删除。",
        "状态字段使用 VARCHAR(32) 保存可读枚举；关键查询必须有组合索引；禁止在高频路径使用 SELECT 后再 UPDATE 的非原子名额逻辑。",
        "数据库变更只能通过 Flyway 版本脚本执行，禁止手工修改生产表结构。",
    ])
    heading(doc, "2. 核心表与关键约束")
    table(doc, ["表", "关键字段", "关键索引或约束"], [
        ["user", "id、login_name、password_hash、phone、email、role、status。", "UK(login_name)、UK(phone，可空策略需统一)、idx(role,status)。"],
        ["organization", "id、name、status、owner_user_id。", "UK(name)、idx(owner_user_id)。"],
        ["activity", "id、organization_id、title、status、registration_start_at、registration_end_at、start_at、end_at。", "idx(organization_id,status)、idx(status,registration_start_at,registration_end_at)。"],
        ["activity_session", "id、activity_id、name、venue、start_at、end_at、status。", "idx(activity_id,start_at)、idx(status,start_at)。"],
        ["activity_quota", "id、session_id、total_quota、available_quota、reserved_quota、confirmed_quota、version。", "UK(session_id)，version 用于后台调整时的乐观锁。"],
        ["registration", "id、user_id、activity_id、session_id、registration_no、status、reservation_expire_at、confirmed_at、cancelled_at、request_id。", "UK(user_id,session_id)、UK(request_id)、UK(registration_no)、idx(session_id,status,create_time)。"],
        ["waitlist", "id、user_id、session_id、status、queue_no、reserved_until、promoted_registration_id。", "UK(user_id,session_id)、UK(session_id,queue_no)、idx(session_id,status,queue_no)。"],
        ["check_in", "id、registration_id、check_in_at、operator_user_id、channel。", "UK(registration_id)、idx(check_in_at)。"],
        ["outbox_event", "id、event_type、aggregate_type、aggregate_id、payload_json、status、retry_count、next_retry_at。", "UK(id)、idx(status,next_retry_at)、idx(aggregate_type,aggregate_id)。"],
        ["operation_log", "id、operator_user_id、operation_type、resource_type、resource_id、request_id、detail_json。", "idx(resource_type,resource_id)、idx(operator_user_id,create_time)、idx(request_id)。"],
    ], [1550, 4450, 3360])
    heading(doc, "3. 名额守恒不变量")
    paragraph(doc, "对每个场次，total_quota = available_quota + reserved_quota + confirmed_quota。RESERVED 和 CONFIRMED 均占用名额；WAITING、CANCELLED、EXPIRED 不占用名额。该不变量是压测、补偿和人工排障的共同校验公式。")
    heading(doc, "4. Redis 键、TTL 与 Lua 原子链路")
    table(doc, ["键", "TTL", "用途"], [
        ["activity:quota:{sessionId}", "活动结束后 24 小时", "保存可用名额和活动是否开放等原子校验信息。"],
        ["registration:lock:{userId}:{sessionId}", "10 分钟或最终状态落库后短暂保留", "防止同一用户重复占位。"],
        ["registration:idempotent:{userId}:{sessionId}:{requestId}", "24 小时", "重复请求返回首次处理结果。"],
        ["registration:status:{registrationNo}", "7 天", "向前端快速返回异步报名处理状态。"],
        ["activity:waitlist:{sessionId}", "活动结束后 24 小时", "ZSET，score 为 queue_no 或进入时间。"],
        ["activity:detail:{activityId}", "10 分钟，发布/下线时主动失效", "活动详情缓存。"],
    ], [3500, 2200, 3660])
    bullets(doc, [
        "Redis Lua 脚本按顺序校验：活动开放、请求幂等、用户占位、可用名额；成功后扣减 available_quota、写入用户锁、记录 registrationNo 和过期时间。",
        "Lua 返回码至少区分：SUCCESS、IDEMPOTENT_REPLAY、ALREADY_REGISTERED、NOT_OPEN、SOLD_OUT、SYSTEM_NOT_READY。",
        "Redis 库存初始化、活动下线、取消、过期和补偿必须走受控服务方法，禁止管理员直接使用 redis-cli 修改名额。",
    ])
    heading(doc, "5. RabbitMQ 与 Outbox")
    table(doc, ["事件", "交换机 / 队列", "消费者动作"], [
        ["registration.reserved", "eventflow.topic / registration.persist", "幂等确认报名占位、更新查询状态、发送站内通知。"],
        ["registration.expire", "eventflow.delay / registration.expire", "校验仍为 RESERVED 且已到期，置 EXPIRED，回补名额并触发递补。"],
        ["quota.released", "eventflow.topic / waitlist.promote", "领取最早有效候补，创建新的 RESERVED 与过期事件。"],
        ["registration.confirmed", "eventflow.topic / notification.registration", "生成通知和动态签到通行证信息。"],
        ["event.failed", "eventflow.dlx / eventflow.dead", "记录失败原因、告警，并等待管理员或补偿任务处置。"],
    ], [2450, 3150, 3760])
    steps(doc, [
        "本地事务写入 registration 状态变化与 outbox_event，不直接在事务中发送 MQ。",
        "Outbox Publisher 查询待投递事件并发布到 RabbitMQ；发布确认成功后标记 PUBLISHED。",
        "消费者先以 eventId 查询消费幂等记录；未处理时在本地事务中执行业务动作和消费记录写入。",
        "失败按有限次数重试；超限进入死信队列；补偿任务基于 Outbox、报名状态和 Redis 键执行对账。",
    ])
    heading(doc, "6. 失败补偿与对账")
    bullets(doc, [
        "每 5 分钟扫描已过期 RESERVED、长时间未投递 Outbox、死信事件和 Redis/MySQL 名额差异。",
        "补偿动作必须幂等并记录 operation_log；不能仅凭 Redis 值覆盖 MySQL 最终事实。",
        "活动开始前和结束后各执行一次全量场次对账，输出可用名额、保留、确认、候补和异常清单。",
    ])
    save(doc, "08-EventFlow数据Redis与消息一致性设计.docx")


def api_auth_spec():
    doc = create_document("EventFlow 接口与权限契约设计", "REST API、JWT、幂等与错误规范 | 技术设计版本 1.0")
    heading(doc, "1. API 通用约定")
    table(doc, ["项目", "约定"], [
        ["基础路径", "/api/v1；资源使用复数名词，URL 使用 kebab-case。"],
        ["认证", "Authorization: Bearer <JWT>；登录、刷新令牌和公开活动浏览接口例外。"],
        ["幂等", "报名、确认、取消和管理员补偿请求必须携带 X-Request-Id；服务端返回相同 requestId 的首次处理结果。"],
        ["时间", "请求与响应采用 ISO-8601 UTC 时间；前端按用户时区展示。"],
        ["分页", "page 从 1 开始，size 默认 20、最大 100；返回 records、page、size、total。"],
        ["响应", "统一结构：code、message、data、requestId、timestamp；成功 code 为 SUCCESS。"],
        ["错误", "使用明确业务 code，例如 REGISTRATION_SOLD_OUT、REGISTRATION_EXPIRED、FORBIDDEN_RESOURCE。"],
    ], [1950, 7410])
    heading(doc, "2. 权限模型")
    table(doc, ["角色", "资源范围", "关键能力"], [
        ["USER", "本人公开可见活动与本人报名。", "浏览、报名、确认、取消、候补查询、二维码和签到。"],
        ["ORGANIZER", "所属 organization 的活动和运营数据。", "创建、编辑、发布、下线、查看名单、管理场次和查看签到。"],
        ["ADMIN", "全平台资源。", "组织者管理、活动治理、全局数据、日志、异常补偿与人工处理。"],
    ], [1700, 3300, 4360])
    paragraph(doc, "授权采用“角色 + 资源归属”双重判断。组织者角色本身不代表可访问所有活动，所有活动读写接口必须校验 activity.organization_id 与当前组织者归属一致。")
    heading(doc, "3. 用户端接口")
    table(doc, ["方法与路径", "用途", "权限与关键规则"], [
        ["POST /auth/login", "账号密码登录。", "公开；返回访问令牌与刷新令牌。"],
        ["POST /auth/refresh", "刷新访问令牌。", "刷新令牌有效且未撤销。"],
        ["GET /activities", "活动列表与筛选。", "公开；仅返回 PUBLISHED 或报名相关可见状态。"],
        ["GET /activities/{activityId}", "活动详情与场次。", "公开；活动下线后按权限返回。"],
        ["POST /sessions/{sessionId}/registrations", "申请报名或进入候补。", "USER；X-Request-Id；调用 Redis Lua。"],
        ["POST /registrations/{registrationNo}/confirm", "确认 RESERVED。", "本人；X-Request-Id；校验 10 分钟有效期。"],
        ["POST /registrations/{registrationNo}/cancel", "取消报名或候补。", "本人；活动开始前 24 小时；X-Request-Id。"],
        ["GET /me/registrations", "我的报名记录。", "本人；支持状态和时间筛选。"],
        ["GET /registrations/{registrationNo}/pass", "获取动态签到二维码令牌。", "本人；仅 CONFIRMED；短时有效。"],
    ], [3300, 2600, 3460])
    heading(doc, "4. 组织者与管理员接口")
    table(doc, ["方法与路径", "用途", "权限"], [
        ["POST /organizer/activities", "创建活动草稿。", "ORGANIZER。"],
        ["PUT /organizer/activities/{activityId}", "编辑草稿或允许编辑的活动信息。", "资源归属组织者。"],
        ["POST /organizer/activities/{activityId}/sessions", "新增或编辑场次与初始名额。", "资源归属组织者；校验时间与容量。"],
        ["POST /organizer/activities/{activityId}/publish", "发布活动并初始化 Redis 名额。", "资源归属组织者；所有发布校验通过。"],
        ["POST /organizer/activities/{activityId}/offline", "组织者下线活动。", "资源归属组织者；记录审计。"],
        ["GET /organizer/activities/{activityId}/registrations", "查询报名名单和状态。", "资源归属组织者；支持导出异步任务。"],
        ["POST /check-ins/verify", "验证二维码并签到。", "ORGANIZER 或授权工作人员；幂等。"],
        ["GET /admin/operations/anomalies", "查询死信、对账差异与异常占位。", "ADMIN。"],
        ["POST /admin/operations/compensations/{id}", "执行人工补偿。", "ADMIN；X-Request-Id；二次确认与审计。"],
    ], [3450, 2920, 2990])
    heading(doc, "5. 报名接口状态语义")
    table(doc, ["结果 code", "HTTP", "前端动作"], [
        ["REGISTRATION_RESERVED", "202", "显示 10 分钟倒计时并进入确认页。"],
        ["REGISTRATION_CONFIRMED", "200", "展示报名成功和动态通行证入口。"],
        ["REGISTRATION_WAITLISTED", "202", "展示候补序号和候补说明。"],
        ["REGISTRATION_ALREADY_EXISTS", "200", "跳转既有报名详情，不重复创建。"],
        ["REGISTRATION_SOLD_OUT", "409", "提示满额并提供加入候补操作。"],
        ["REGISTRATION_EXPIRED", "409", "提示保留已过期并刷新场次状态。"],
        ["REGISTRATION_CANCELLATION_CLOSED", "409", "说明已超过活动前 24 小时取消窗口。"],
    ], [3300, 1300, 3760])
    heading(doc, "6. 接口安全与可观测性")
    bullets(doc, [
        "对登录、报名、确认、取消和签到接口实施基于用户与 IP 的限流；限流返回明确 code 与可重试提示。",
        "OpenAPI 文档仅暴露在受保护的测试或管理环境；生产环境不展示敏感调试信息。",
        "手机号、邮箱在组织者列表默认脱敏；只有经过授权的导出任务才可获取完整数据。",
        "所有写接口写入 requestId；管理员和组织者的关键写操作写入 operation_log。",
    ])
    save(doc, "09-EventFlow接口与权限契约设计.docx")


def main():
    architecture_spec()
    data_event_spec()
    api_auth_spec()


if __name__ == "__main__":
    main()
