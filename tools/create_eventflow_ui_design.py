from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("docs/project/06-EventFlowUI与交互设计规范.docx")
ASSETS = Path("assets/ui-concepts/v4")
CONTENT_WIDTH_DXA = 9360
BLUE = RGBColor(0, 229, 238)
NAVY = RGBColor(10, 26, 32)
GRAY = RGBColor(93, 111, 118)


def set_font(run, size=11, color=None, bold=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_width(cell, width):
    properties = cell._tc.get_or_add_tcPr()
    value = properties.first_child_found_in("w:tcW")
    if value is None:
        value = OxmlElement("w:tcW")
        properties.append(value)
    value.set(qn("w:w"), str(width))
    value.set(qn("w:type"), "dxa")


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(0, 134, 145), 18, 10),
        ("Heading 2", 13, RGBColor(0, 134, 145), 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("EventFlow UI Design Specification")
    set_font(run, 8.5, GRAY)


def heading(doc, text, level=1):
    doc.add_paragraph(text, style=f"Heading {level}")


def paragraph(doc, text, bold_prefix=None):
    item = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        prefix = item.add_run(bold_prefix)
        set_font(prefix, 11, bold=True)
        suffix = item.add_run(text[len(bold_prefix):])
        set_font(suffix)
    else:
        run = item.add_run(text)
        set_font(run)


def bullets(doc, values):
    for value in values:
        item = doc.add_paragraph(style="List Bullet")
        item.paragraph_format.space_after = Pt(4)
        run = item.add_run(value)
        set_font(run)


def table(doc, headers, rows, widths):
    value = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        cell = value.rows[0].cells[index]
        shade(cell, "DDF4F4")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        item = cell.paragraphs[0]
        item.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = item.add_run(header)
        set_font(run, 9.5, bold=True)
    for row in rows:
        cells = value.add_row().cells
        for index, text in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            run = cell.paragraphs[0].add_run(text)
            set_font(run, 9.5)
    for row in value.rows:
        for index, cell in enumerate(row.cells):
            set_width(cell, widths[index])
            for item in cell.paragraphs:
                item.paragraph_format.space_after = Pt(3)
                item.paragraph_format.line_spacing = 1.15
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def concept(doc, filename, caption):
    path = ASSETS / filename
    doc.add_picture(str(path), width=Inches(6.35))
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(10)
    run = label.add_run(caption)
    set_font(run, 9, GRAY)


def main():
    doc = Document()
    configure_document(doc)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("EventFlow UI 与交互设计规范")
    set_font(run, 22, NAVY, True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("设计方向：未来文明活动入口 | UI 概念版本 4.0")
    set_font(run, 10.5, GRAY)

    heading(doc, "1. 设计定位")
    paragraph(doc, "EventFlow 采用“未来文明活动入口”视觉方向：界面本身是未来系统，而不是被放在电脑模型中的截图。用户首先进入巨型光环、悬浮舞台与粒子潮汐构成的活动世界，再完成报名；组织者和管理员则在空间化名额编排界面中操作。科技感来自珠光白操作面、深海蓝核心场景、电青能量路径和建筑级空间叙事，而不是会议人群、普通后台卡片或游戏化 HUD。")
    table(doc, ["设计原则", "落地规则"], [
        ["全屏直接感", "概念图和正式页面都不使用电脑、手机或浏览器外框；首屏直接呈现产品界面和关键业务动作。"],
        ["电影级叙事", "用户端以巨型光环、悬浮舞台和活动世界为视觉焦点；报名按钮、名额和状态必须在首屏清晰可见。"],
        ["空间化科技", "珠光白操作面与深海蓝核心场景形成反差，电青承载信息路径，暖橙仅表示时间压力和告警。"],
        ["运营优先", "后台使用固定侧边导航、筛选工具栏、数据表格、状态标签和实时数据面板。"],
        ["状态可感知", "名额锁定、确认、候补、过期和签到必须通过文字、颜色、图标和动作反馈共同表达。"],
    ], [2400, 6960])

    heading(doc, "2. 视觉系统")
    table(doc, ["Token", "建议值", "用途"], [
        ["背景", "#F4F7F9 / #E9F0F4", "珠光白页面基础层与浅灰蓝建筑线条。"],
        ["核心场景", "#051B3A / #0A2850", "活动光门、名额反应器和关键控制面。"],
        ["表面", "#FFFFFF / #EDF5F7", "表格、表单和有限的操作面板。"],
        ["主色", "#00D7E8", "主操作、焦点、容量进度、激活导航与能量路径。"],
        ["强调色", "#FF9D2E", "名额紧张、保留倒计时、异常和告警。"],
        ["成功", "#39D98A", "报名确认、签到成功、服务正常。"],
        ["文本", "#0B214A / #5C7280", "深海军蓝主文字与次级说明。"],
        ["圆角", "4px 至 8px", "工具化界面，避免营销式大圆角卡片。"],
        ["动效", "160ms 至 240ms", "按钮反馈、状态切换、数字刷新；不使用长时间装饰动画。"],
    ], [1750, 2200, 5410])
    bullets(doc, [
        "图标统一使用细线型图标；实际前端优先采用 Lucide 图标库。",
        "正文保持高可读性无衬线字体；容量、时间、倒计时和编号可使用等宽数字样式。",
        "所有科技纹理只用于背景层，文本与表格区域保持高对比、低噪声。",
    ])

    heading(doc, "3. 页面架构")
    table(doc, ["用户端", "组织者与管理员端"], [
        ["登录；活动列表；活动详情；场次选择；报名确认；我的报名；报名详情；动态签到二维码。", "运营总览；活动管理；活动创建与编辑；场次与名额；报名名单；候补队列；签到记录；活动数据统计；用户与组织者管理；操作日志。"],
    ], [4680, 4680])
    paragraph(doc, "导航规则：用户端使用轻量顶栏和账户菜单；运营端使用固定左侧导航与顶部上下文栏。管理员端在组织者菜单基础上增加用户、审核、日志和异常补偿入口。")

    heading(doc, "4. 核心屏幕概念")
    heading(doc, "4.1 用户活动探索与报名", 2)
    concept(doc, "v4-cinematic-event-entry.png", "概念图 1：未来文明活动入口，以巨型光环、粒子潮汐、名额信号与场次时间线建立叙事。")
    bullets(doc, [
        "首屏按视觉优先级展示：活动封面、活动名称、时间地点、名额状态、报名按钮、场次选择。",
        "名额数字使用文字加进度条，不能只以颜色表达剩余容量。",
        "报名未开始、报名中、名额紧张、满额候补、报名结束均需要明确按钮和状态文案。",
    ])

    heading(doc, "4.2 报名确认与临时保留", 2)
    concept(doc, "v4-cinematic-reservation-flow.png", "概念图 2：名额光舱与动态通行证交互，用户从场次选择、身份信息到确认进入活动世界。")
    bullets(doc, [
        "以三步进度表示场次选择、报名信息与确认结果；用户始终知道自己当前所在阶段。",
        "倒计时使用暖橙强调，并同时显示“保留至具体时间”，避免仅依赖动态数字。",
        "确认按钮必须防重复点击；提交中显示 loading；失败后保留用户输入并给出可执行的错误提示。",
    ])

    heading(doc, "4.3 活动运营指挥台", 2)
    concept(doc, "v4-cinematic-quota-command.png", "概念图 3：空间化名额指挥面，将配额、候补、签到、活动阶段与服务状态组织为单一编排系统。")
    bullets(doc, [
        "仪表盘只呈现当前活动最需要行动的数据：容量、报名趋势、候补、签到和异常。",
        "表格支持按场次、状态、时间和关键词筛选；筛选条件在页面顶部固定可见。",
        "名额告急、死信或补偿异常使用暖橙或红色，并提供可进入详情的明确操作。",
    ])

    heading(doc, "4.4 活动创建向导", 2)
    concept(doc, "v4-cinematic-activity-builder.png", "概念图 4：活动建筑师控制面，以事件配置、场次轨道和配额蓝图完成创建与发布。")
    bullets(doc, [
        "创建流程采用左侧步骤导航：活动信息、场次、名额、报名窗口、预览发布。",
        "草稿保存始终可用；发布前统一校验必填项、时间关系、场次重叠和名额合法性。",
        "场次和名额使用可编辑表格，避免在多个模态框中完成核心配置。",
    ])

    heading(doc, "5. 组件与状态规范")
    table(doc, ["组件", "设计与交互规则"], [
        ["按钮", "主操作使用电青实色；次操作使用描边；删除、下线和异常补偿使用文字加图标并二次确认。"],
        ["状态标签", "RESERVED 使用暖橙；CONFIRMED 和 CHECKED_IN 使用绿色；WAITING 使用蓝色；CANCELLED、EXPIRED 使用灰色。"],
        ["数据表格", "表头固定、行高稳定、短字段居中、说明字段左对齐；操作采用图标按钮并提供 tooltip。"],
        ["表单", "输入框使用白色表面与浅蓝边界；焦点电青描边；必填、错误、禁用和 loading 状态均独立表达。"],
        ["倒计时", "仅在 RESERVED 或候补递补保留中展示；时间到期后自动刷新状态而非让用户继续提交。"],
        ["空状态与错误", "使用简洁图标和下一步操作；不使用装饰性大插画遮蔽问题原因。"],
    ], [2100, 7260])

    heading(doc, "6. 响应式与可访问性")
    bullets(doc, [
        "桌面端以 1440px 设计基准，后台内容最小宽度以表格可读性为优先。",
        "移动端用户页面保留活动浏览、报名、我的报名和二维码；运营后台以关键查看和轻操作为主。",
        "移动端导航收为抽屉；场次、候补和报名记录改为纵向列表，不压缩成不可读的表格。",
        "文字与背景至少保持 WCAG AA 对比度；颜色不作为唯一状态信息；所有图标按钮提供可访问名称。",
        "键盘可完成登录、筛选、表单填写、确认报名和关闭弹层等关键操作。",
    ])

    heading(doc, "7. 核心交互语言")
    table(doc, ["业务动作", "界面交互与反馈"], [
        ["选择场次", "用户点击时间线或场次卡后，容量信号平滑切换，按钮文案与剩余名额同步更新；已满场次提供“加入候补”而不是禁用后无解释。"],
        ["锁定名额", "提交报名后进入短暂的信号扫描 loading，成功后切换到 RESERVED，并固定显示 10 分钟倒计时和明确确认入口。"],
        ["确认报名", "确认按钮在请求中不可重复触发；成功后呈现动态通行证和签到二维码入口，失败后保留表单内容并说明可重试原因。"],
        ["名额实时变化", "容量圆环、进度条和数据表只更新变化字段，使用 160ms 至 240ms 的数字过渡，避免整屏闪烁。"],
        ["候补递补", "候补转 RESERVED 时使用暖橙高优先级通知，并在报名详情中展示新的保留截止时间和确认操作。"],
        ["异常处理", "管理员的下线、补偿和删除操作必须有二次确认；执行结果以操作日志与非阻塞通知共同反馈。"],
    ], [2250, 7110])

    heading(doc, "8. 资产清单与实现边界")
    table(doc, ["资产", "位置", "用途"], [
        ["用户端概念图", "assets/ui-concepts/v4/v4-cinematic-event-entry.png", "未来活动入口、名额信号与场次时间线参考。"],
        ["报名确认概念图", "assets/ui-concepts/v4/v4-cinematic-reservation-flow.png", "名额光舱、锁定倒计时和动态通行证参考。"],
        ["运营端概念图", "assets/ui-concepts/v4/v4-cinematic-quota-command.png", "空间化配额编排、候补与签到指挥面参考。"],
        ["活动创建概念图", "assets/ui-concepts/v4/v4-cinematic-activity-builder.png", "活动建筑师控制面、配额与时间线配置参考。"],
    ], [2350, 4550, 2460])
    paragraph(doc, "实现边界：概念图用于确定布局、层级、色彩和组件语言。实际页面的中文内容、按钮、表格、图表、表单和状态逻辑必须由前端代码实现，不能将截图中的文字或数据直接当作页面素材。")

    heading(doc, "9. 界面语言规范")
    table(doc, ["规则", "正式实现要求"], [
        ["默认语言", "用户端、组织者端和管理员端的可见业务文字默认使用简体中文。"],
        ["产品名称", "产品品牌统一写作 EventFlow；首次出现时可搭配“活动报名与名额预约平台”。"],
        ["英文保留范围", "仅保留 API、JWT、Redis、RabbitMQ、ID、URL、QR Code 等具有明确技术或行业含义的英文缩写。"],
        ["业务状态", "页面展示使用中文状态，例如“名额锁定中”“报名已确认”“候补中”“已过期”“已签到”；代码枚举仍可使用 RESERVED、CONFIRMED 等英文值。"],
        ["按钮与表单", "使用“立即报名”“确认报名”“加入候补”“保存草稿”“发布活动”“查看详情”等直接中文动作词，不使用无必要的英文按钮。"],
        ["概念图文字", "当前生成图中的英文仅用于布局、排版和信息层级参考，正式前端不得直接作为界面文案。"],
    ], [2250, 7110])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
